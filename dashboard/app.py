"""
GeoAlert Command Hub v3
=======================
Landslide early-warning command centre (Streamlit + Plotly + Firebase RTDB).

Highlights
  * Alert debouncing (hysteresis) + optional webhook notifications
  * Infinite-slope Factor of Safety with Monte-Carlo reliability (probability of failure)
  * Fan-chart forecasts with breach probabilities, multivariate anomaly detection
  * Multi-node fleet view + map, incident replay, node-health / sensor-QC checks
  * Plain-language auto-generated situation report
  * Exports: Excel, CSV, JSON, GeoJSON, interactive HTML, PDF, and a one-click ZIP bundle

Requires: streamlit>=1.37 plotly pandas numpy requests openpyxl
Optional: reportlab matplotlib (PDF report)

Run: streamlit run geoalert_hub_v3.py
"""
from __future__ import annotations

import inspect
import io
import json
import time
import zipfile
import zlib
from datetime import datetime, timezone
from types import SimpleNamespace

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import requests
import streamlit as st
from plotly.subplots import make_subplots

# ==============================================================================
# 1. PAGE SETUP & THEME
# ==============================================================================
st.set_page_config(page_title="GeoAlert Command Hub", page_icon="🛡️", layout="wide", initial_sidebar_state="expanded")

# Streamlit >= 1.50 replaced use_container_width with width="stretch"
NEW_API = "width" in inspect.signature(st.plotly_chart).parameters
WIDE = {"width": "stretch"} if NEW_API else {"use_container_width": True}

st.markdown(
    """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;800&display=swap');
    html, body, .stApp { font-family: 'Inter', 'Segoe UI', sans-serif; }
    .stApp { background: radial-gradient(1200px 600px at 10% -10%, #14213D 0%, #0B1120 55%); color: #F1F5F9; }
    div[data-testid="stMetric"] {
        background: rgba(30,41,59,.65); backdrop-filter: blur(6px);
        border: 1px solid rgba(148,163,184,.15); border-radius: 12px; padding: 14px;
        box-shadow: 0 6px 16px -6px rgba(0,0,0,.5);
    }
    div[data-testid="stMetricLabel"] { color:#94A3B8 !important; font-weight:600; font-size:.78rem;
        letter-spacing:.06em; text-transform:uppercase; }
    div[data-testid="stMetricValue"] { color:#F8FAFC !important; font-weight:800; font-size:1.55rem; }
    .status-card { padding:20px 24px; border-radius:14px; margin-bottom:18px; border:1px solid rgba(255,255,255,.08); }
    .brief { background: rgba(30,41,59,.6); border:1px solid rgba(148,163,184,.18); border-left:5px solid #38BDF8;
             border-radius:10px; padding:14px 18px; line-height:1.55; margin:6px 0 14px 0; }
    .chip { display:inline-block; padding:2px 10px; border-radius:999px; font-size:.72rem; font-weight:700;
            margin-right:6px; background:rgba(148,163,184,.18); color:#E2E8F0; }
    .pulse { animation: pulse 1.3s infinite; }
    @keyframes pulse { 0%{opacity:1} 50%{opacity:.4} 100%{opacity:1} }
</style>
""",
    unsafe_allow_html=True,
)

# ==============================================================================
# 2. CONFIG, SECRETS, SIDEBAR (part 1: connection)
# ==============================================================================
DEFAULT_DB_URL = "https://geoalert-ai-ee1fd-default-rtdb.asia-southeast1.firebasedatabase.app"
DEFAULT_NODE = "geoalert-node-01"


def _secret(key: str) -> str:
    try:
        return st.secrets.get(key, "")
    except Exception:
        return ""


API_KEY, EMAIL, PASSWORD = _secret("FIREBASE_API_KEY"), _secret("FIREBASE_EMAIL"), _secret("FIREBASE_PASSWORD")
WEBHOOK_SECRET = _secret("ALERT_WEBHOOK_URL")

for _k, _v in (("alert_log", []), ("last_tier", None), ("last_anom_ts", None), ("exports", None)):
    st.session_state.setdefault(_k, _v)

with st.sidebar:
    st.image("https://img.icons8.com/fluency/96/mountain.png", width=64)
    st.title("GeoAlert Hub")
    st.caption("AI-Powered Landslide Disaster Intelligence")
    st.divider()
    st.subheader("⚙️ Connection")
    db_url = st.text_input("Database URL", value=DEFAULT_DB_URL)
    if not API_KEY or not EMAIL:
        st.warning("Firebase credentials not found in secrets. Configure below:")
        API_KEY = st.text_input("Firebase API Key", value=API_KEY, type="password")
        EMAIL = st.text_input("Auth Email", value=EMAIL)
        PASSWORD = st.text_input("Auth Password", value=PASSWORD, type="password")


# ==============================================================================
# 3. FIREBASE REST PIPELINE
# ==============================================================================
def get_token():
    s = st.session_state
    if s.get("fb_token") and time.time() < s.get("fb_token_exp", 0):
        return s["fb_token"]
    if not (API_KEY and EMAIL and PASSWORD):
        return None
    url = f"https://identitytoolkit.googleapis.com/v1/accounts:signInWithPassword?key={API_KEY}"
    try:
        r = requests.post(url, json={"email": EMAIL, "password": PASSWORD, "returnSecureToken": True}, timeout=8)
        r.raise_for_status()
        data = r.json()
        s["fb_token"] = data.get("idToken")
        s["fb_token_exp"] = time.time() + int(data.get("expiresIn", 3600)) - 120
        s["auth_error"] = None
        return s["fb_token"]
    except Exception as err:
        s["auth_error"] = str(err)
        return None


def fb_get(path: str, params: dict | None = None):
    p = dict(params or {})
    token = get_token()
    if token:
        p["auth"] = token
    try:
        r = requests.get(f"{db_url.rstrip('/')}/{path}.json", params=p, timeout=7)
        r.raise_for_status()
        return r.json()
    except Exception:
        return None



def list_nodes():
    s = st.session_state
    if s.get("nodes") is not None and time.time() - s.get("nodes_ts", 0) < 60:
        return s["nodes"]
    data = fb_get("nodes", {"shallow": "true"})
    nodes = sorted(data.keys()) if isinstance(data, dict) and data else []
    s["nodes"], s["nodes_ts"] = nodes, time.time()
    return nodes


# ==============================================================================
# 2b. SIDEBAR (part 2: node, alerting, model parameters)
# ==============================================================================
with st.sidebar:
    _nodes = list_nodes()
    _choices = _nodes if _nodes else [DEFAULT_NODE]
    NODE_ID = st.selectbox("Monitored node", _choices, index=_choices.index(DEFAULT_NODE) if DEFAULT_NODE in _choices else 0)
    auto_refresh = st.checkbox("Auto refresh", value=True)
    refresh_s = st.slider("Refresh interval (s)", 2, 30, 5, disabled=not auto_refresh)
    FC_H = st.slider("Forecast horizon (min)", 15, 180, 60, 15)

    with st.expander("🎚️ Risk thresholds"):
        TH = {
            "tilt_warn": st.number_input("Tilt - warning (°)", 1.0, 45.0, 7.0, 0.5),
            "tilt_high": st.number_input("Tilt - high risk (°)", 1.0, 45.0, 14.0, 0.5),
            "tilt_crit": st.number_input("Tilt - critical (°)", 1.0, 45.0, 25.0, 0.5),
            "moist_crit": st.number_input("Saturation - critical (%)", 50.0, 100.0, 85.0, 1.0),
        }
    with st.expander("🪨 Soil / slope model"):
        SOIL = {
            "c": st.number_input("Effective cohesion c' (kPa)", 0.0, 50.0, 5.0, 0.5),
            "phi": st.number_input("Friction angle phi' (°)", 10.0, 45.0, 30.0, 0.5),
            "gamma": st.number_input("Soil unit weight (kN/m³)", 12.0, 24.0, 18.0, 0.5),
            "depth": st.number_input("Failure-plane depth z (m)", 0.5, 10.0, 2.0, 0.5),
        }
        CV = {
            "c": st.slider("Cohesion uncertainty (CV %)", 0, 60, 25) / 100,
            "phi": st.slider("Friction uncertainty (CV %)", 0, 30, 8) / 100,
            "gamma": 0.05,
        }
    with st.expander("🔔 Alerting"):
        DEBOUNCE = st.checkbox("Debounce tier changes", True,
                               help="Escalate only after N consecutive samples; de-escalate only after M. CRITICAL is never delayed.")
        K_UP = st.slider("Samples to escalate (N)", 1, 8, 3, disabled=not DEBOUNCE)
        K_DOWN = st.slider("Samples to de-escalate (M)", 3, 30, 10, disabled=not DEBOUNCE)
        NOTIFY_ON = st.checkbox("Send webhook on escalation", False)
        WEBHOOK = WEBHOOK_SECRET or st.text_input("Webhook URL (Slack/Discord/custom)", type="password")
        NOTIFY_MIN = {"WARNING": 1, "HIGH RISK": 2, "CRITICAL": 3}[st.selectbox("Notify from tier", ["WARNING", "HIGH RISK", "CRITICAL"], index=1)]
        NOTIFY_COOLDOWN = st.slider("Reminder cooldown (min)", 1, 60, 10)
    with st.expander("🗺️ Site coordinates"):
        st.caption("Used for nodes without `nodes/<id>/meta {lat, lon}` (demo positions are offset around this point).")
        BASE_LAT = st.number_input("Latitude", -90.0, 90.0, 13.0, 0.001, format="%.4f")
        BASE_LON = st.number_input("Longitude", -180.0, 180.0, 75.0, 0.001, format="%.4f")
    st.divider()
    st.markdown("📍 **Sector:** Western Ridge Zone 4\n\n🛰️ **Hardware:** ESP32 + SIM800L Dual-Uplink")

# ==============================================================================
# 4. GEOTECHNICAL & ANALYTICS ENGINE
# ==============================================================================
TIERS = {
    "NORMAL": {"color": "#10B981", "level": 0},
    "WARNING": {"color": "#F59E0B", "level": 1},
    "HIGH RISK": {"color": "#F97316", "level": 2},
    "CRITICAL": {"color": "#EF4444", "level": 3},
}
LEVEL_NAME = {v["level"]: k for k, v in TIERS.items()}

PROTOCOL = {
    "NORMAL": ["Continue routine telemetry monitoring.", "Verify solar charge and uplink health."],
    "WARNING": ["Increase sampling frequency and watch trend slopes.", "Notify the field engineer on duty.",
                "Inspect drainage channels upslope."],
    "HIGH RISK": ["Alert the local disaster-management authority.", "Pre-position response teams; restrict access to the runout zone.",
                  "Prepare the SMS broadcast to registered residents."],
    "CRITICAL": ["Dispatch configured emergency SMS/voice alerts and follow the approved evacuation protocol.", "Evacuate lower runoff valleys and close roads below the slope.",
                 "Escalate to the district emergency operations centre."],
}
DRIVER_LABEL = {"moisture": "soil saturation", "rain": "rainfall", "tilt": "slope tilt", "tremor": "ground tremor"}
SENSOR_COLS = ["soil_moisture_pct", "rainfall_pct", "temperature_c", "humidity_pct",
               "tilt_angle_deg", "pitch_deg", "roll_deg", "yaw_deg", "vibration_pulses"]
EXPORT_COLS = ["timestamp"] + SENSOR_COLS + ["antecedent_rain", "threat_score", "fos", "risk_tier"]
DARK = dict(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="#0F172A")


def threat_score(moisture, rain, tilt, vibration):
    """Weighted 0-100 index of pore-water pressure, rainfall, displacement and tremor."""
    return (np.minimum(moisture, 100) * 0.35 + np.minimum(rain, 100) * 0.25
            + np.minimum(np.asarray(tilt) / 45.0 * 100, 100) * 0.30 + np.minimum(np.asarray(vibration) * 10, 100) * 0.10)


def score_drivers(s):
    return {"moisture": min(s["soil_moisture_pct"], 100) * 0.35, "rain": min(s["rainfall_pct"], 100) * 0.25,
            "tilt": min(s["tilt_angle_deg"] / 45 * 100, 100) * 0.30, "tremor": min(s["vibration_pulses"] * 10, 100) * 0.10}


def factor_of_safety(tilt_deg, moisture_pct, soil):
    """Infinite slope with pore pressure: (c' + (gz cos²b - u) tan(phi')) / (gz sin b cos b). Arrays broadcast."""
    beta = np.radians(np.clip(tilt_deg, 1.0, 60.0))
    m = np.clip(moisture_pct, 0, 100) / 100.0
    gz = soil["gamma"] * soil["depth"]
    cos2 = np.cos(beta) ** 2
    u = m * 9.81 * soil["depth"] * cos2
    num = soil["c"] + (gz * cos2 - u) * np.tan(np.radians(soil["phi"]))
    return np.clip(num / (gz * np.sin(beta) * np.cos(beta)), 0, 5)


def classify(score, tilt, moisture, rain, fos, th):
    if score >= 75 or tilt >= th["tilt_crit"] or (moisture > th["moist_crit"] and rain > 70) or fos < 1.0:
        name = "CRITICAL"
    elif score >= 55 or tilt >= th["tilt_high"] or moisture >= 75 or fos < 1.25:
        name = "HIGH RISK"
    elif score >= 35 or tilt >= th["tilt_warn"] or moisture >= 50:
        name = "WARNING"
    else:
        name = "NORMAL"
    return {"name": name, **TIERS[name], "score": float(score), "fos": float(fos)}


def reliability(tilt, moisture, soil, cv, n=4000, seed=3):
    """Monte-Carlo probability of failure P(FoS < 1) under soil-parameter and moisture-sensor uncertainty."""
    rng = np.random.default_rng(seed)
    ss = {
        "c": np.clip(rng.normal(soil["c"], max(soil["c"] * cv["c"], 1e-9), n), 0, None),
        "phi": np.clip(rng.normal(soil["phi"], soil["phi"] * cv["phi"], n), 5, 50),
        "gamma": np.clip(rng.normal(soil["gamma"], soil["gamma"] * cv["gamma"], n), 10.5, None),
        "depth": soil["depth"],
    }
    f = factor_of_safety(tilt, np.clip(rng.normal(moisture, 3.0, n), 0, 100), ss)
    sd = float(f.std())
    return {"f": f, "pf": float((f < 1).mean()), "mean": float(f.mean()), "sd": sd,
            "beta": float((f.mean() - 1) / sd) if sd > 1e-9 else float("inf")}


def forecast_eta(df, col, limit, window=20):
    d = df.tail(window).dropna(subset=[col])
    if len(d) < 5:
        return None
    x = (d["timestamp"] - d["timestamp"].iloc[0]).dt.total_seconds().to_numpy() / 60.0
    if np.ptp(x) == 0:
        return None
    slope, icpt = np.polyfit(x, d[col].to_numpy(), 1)
    cur = float(slope * x[-1] + icpt)
    eta = 0.0 if cur >= limit else (None if slope <= 1e-6 else (limit - cur) / slope)
    return {"slope": float(slope), "cur": cur, "eta_min": eta}


def fmt_eta(fc):
    if fc is None:
        return "insufficient data"
    m = fc["eta_min"]
    if m is None:
        return "not rising"
    if m == 0:
        return "limit reached"
    if m > 24 * 60:
        return "over 24 h"
    return f"{m:.0f} min" if m < 120 else f"{m / 60:.1f} h"


def mc_forecast(df, col, horizon, window=40, sims=400, seed=11):
    """Linear-trend Monte-Carlo (slope uncertainty + residual noise) -> percentile fan + raw paths."""
    d = df.tail(window).dropna(subset=[col])
    if len(d) < 8:
        return None
    x = (d["timestamp"] - d["timestamp"].iloc[0]).dt.total_seconds().to_numpy() / 60.0
    if np.ptp(x) <= 0:
        return None
    y = d[col].to_numpy(float)
    slope, icpt = np.polyfit(x, y, 1)
    fit = slope * x + icpt
    s = float(np.sqrt(((y - fit) ** 2).sum() / max(len(x) - 2, 1)))
    sxx = float(((x - x.mean()) ** 2).sum())
    se = s / np.sqrt(sxx) if sxx > 0 else 0.0
    steps = int(min(horizon, 120))
    t = np.arange(1, steps + 1) * (horizon / steps)
    rng = np.random.default_rng(seed)
    paths = np.maximum(fit[-1] + rng.normal(slope, se, (sims, 1)) * t[None, :] + rng.normal(0, s, (sims, steps)), 0)
    return {"t": t, "p10": np.percentile(paths, 10, 0), "p50": np.percentile(paths, 50, 0),
            "p90": np.percentile(paths, 90, 0), "paths": paths, "slope": float(slope), "cur": float(fit[-1]), "sigma": s}


def p_breach(mc, limit):
    return None if mc is None else float((mc["paths"].max(axis=1) >= limit).mean())


def flag_anomalies(series: pd.Series, window=12, z=3.0) -> pd.Series:
    d = series.diff()
    mu, sd = d.rolling(window, min_periods=5).mean(), d.rolling(window, min_periods=5).std().replace(0, np.nan)
    return ((d - mu).abs() > z * sd).fillna(False)


def mahalanobis_scores(df, cols):
    """Robust multivariate anomaly score (median/MAD standardisation + Mahalanobis distance)."""
    X = df[cols].to_numpy(float)
    med = np.median(X, axis=0)
    mad = np.median(np.abs(X - med), axis=0) * 1.4826
    sd = np.where(mad < 1e-9, X.std(axis=0), mad)
    sd = np.where(sd < 1e-9, 1.0, sd)
    Z = (X - med) / sd
    inv = np.linalg.pinv(np.cov(Z, rowvar=False) + np.eye(len(cols)) * 1e-3)
    return np.sqrt(np.maximum(np.einsum("ij,jk,ik->i", Z, inv, Z), 0))


def chi_threshold(k, z=2.326):
    """sqrt of the 99th-percentile chi-square (Wilson-Hilferty approximation)."""
    return float(np.sqrt(k * (1 - 2 / (9 * k) + z * np.sqrt(2 / (9 * k))) ** 3))


def log_event(level: str, message: str):
    st.session_state.alert_log.insert(0, {"time_utc": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
                                          "level": level, "event": message, "ack": False})
    del st.session_state.alert_log[200:]


# ==============================================================================
# 5. DATA PARSING, HISTORY, FLEET, HEALTH
# ==============================================================================
def _f(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def parse_live(live):
    if not live:
        s = {"soil_moisture_pct": 68.4, "rainfall_pct": 54.0, "temperature_c": 23.5, "humidity_pct": 89.0,
             "tilt_angle_deg": 12.8, "pitch_deg": 11.4, "roll_deg": 5.8, "yaw_deg": 1.2, "vibration_pulses": 4}
        return s, {"active_interface": "DEMO / RECONNECTING", "rssi": -68}, False, {}
    raw = live.get("sensor_readings", live)
    mpu = raw.get("mpu_raw", {}) or {}
    s = {"soil_moisture_pct": _f(raw.get("soil_moisture_pct")), "rainfall_pct": _f(raw.get("rainfall_pct")),
         "temperature_c": _f(raw.get("temperature_c")), "humidity_pct": _f(raw.get("humidity_pct")),
         "tilt_angle_deg": _f(raw.get("tilt_angle_deg")),
         "pitch_deg": _f(mpu.get("pitch_deg", raw.get("pitch_deg"))), "roll_deg": _f(mpu.get("roll_deg", raw.get("roll_deg"))),
         "yaw_deg": _f(mpu.get("yaw_deg", raw.get("yaw_deg"))), "vibration_pulses": int(_f(raw.get("vibration_pulses")))}
    power = live.get("power", {}) or {}
    extras = {}
    for k in ("battery_v", "battery_voltage", "battery_pct", "solar_v", "solar_voltage", "uptime_s"):
        for src in (raw, live, power):
            if isinstance(src, dict) and src.get(k) is not None:
                extras[k] = _f(src[k])
                break
    return s, live.get("network_status", {"active_interface": "Wi-Fi (Online)", "rssi": -62}), True, extras


def normalize_history(raw, sensors):
    """Firebase push-id dict -> tidy, time-sorted DataFrame (or None)."""
    if not raw or not isinstance(raw, dict):
        return None
    recs = [v for v in raw.values() if isinstance(v, dict)]
    if not recs:
        return None
    df = pd.json_normalize(recs)
    df.columns = [c.split(".")[-1] for c in df.columns]
    df = df.loc[:, ~df.columns.duplicated()]
    ts_col = next((c for c in ("timestamp", "ts", "time", "datetime") if c in df.columns), None)
    ts = None
    if ts_col is not None:
        s = df[ts_col]
        if pd.api.types.is_numeric_dtype(s):
            ts = pd.to_datetime(s, unit="ms" if s.dropna().median() > 1e11 else "s", utc=True, errors="coerce")
        else:
            ts = pd.to_datetime(s, utc=True, errors="coerce")
        if ts.notna().any() and ts.min().year < 2000:  # device uptime rather than wall-clock
            ts = None
    if ts is None:
        ts = pd.date_range(end=pd.Timestamp.now(tz="UTC"), periods=len(df), freq="min")
    df["timestamp"] = ts
    for c in SENSOR_COLS:
        df[c] = pd.to_numeric(df[c], errors="coerce") if c in df.columns else sensors[c]
        df[c] = df[c].ffill().bfill().fillna(sensors[c])
    df = df.dropna(subset=["timestamp"]).sort_values("timestamp").tail(500).reset_index(drop=True)
    return df[["timestamp"] + SENSOR_COLS] if len(df) else None


def synth_history(sensors, node_id, n=120):
    """Deterministic demo history ending exactly at the current reading."""
    rng = np.random.default_rng(zlib.crc32(node_id.encode()) & 0xFFFFFFFF)
    t = pd.date_range(end=pd.Timestamp.now(tz="UTC").floor("min"), periods=n, freq="min")
    idx = np.arange(n) - (n - 1)

    def series(end, noise, drift, lo, hi):
        walk = np.cumsum(rng.normal(0, noise, n))
        return np.clip(end + drift * idx + (walk - walk[-1]), lo, hi)

    df = pd.DataFrame({
        "timestamp": t,
        "soil_moisture_pct": series(sensors["soil_moisture_pct"], 0.6, 0.12, 0, 100),
        "rainfall_pct": series(sensors["rainfall_pct"], 1.2, 0.10, 0, 100),
        "temperature_c": series(sensors["temperature_c"], 0.08, -0.01, -10, 60),
        "humidity_pct": series(sensors["humidity_pct"], 0.3, 0.02, 0, 100),
        "tilt_angle_deg": series(sensors["tilt_angle_deg"], 0.12, 0.035, 0, 45),
        "pitch_deg": series(sensors["pitch_deg"], 0.1, 0.03, -45, 45),
        "roll_deg": series(sensors["roll_deg"], 0.1, 0.015, -45, 45),
        "yaw_deg": series(sensors["yaw_deg"], 0.05, 0.0, -180, 180),
    })
    vib = rng.poisson(max(sensors["vibration_pulses"] - 1, 0) + 1, n).astype(float)
    vib[-1] = sensors["vibration_pulses"]
    df["vibration_pulses"] = vib
    return df


def enrich(df, th, soil):
    d = df.copy()
    d["antecedent_rain"] = d["rainfall_pct"].ewm(span=60, adjust=False).mean()
    d["threat_score"] = threat_score(d["soil_moisture_pct"], d["rainfall_pct"], d["tilt_angle_deg"], d["vibration_pulses"])
    d["fos"] = factor_of_safety(d["tilt_angle_deg"], d["soil_moisture_pct"], soil)
    d["risk_tier"] = [classify(s, t, m, r, f, th)["name"] for s, t, m, r, f in
                      zip(d["threat_score"], d["tilt_angle_deg"], d["soil_moisture_pct"], d["rainfall_pct"], d["fos"])]
    return d


def apply_debounce(levels):
    """Hysteresis on the alert level. Returns (stable_level, raw_level)."""
    cur = levels[-1]
    key = f"stable_level_{NODE_ID}"
    prev = st.session_state.get(key)
    if not DEBOUNCE or prev is None or cur == 3:
        new = cur
    elif len(levels) >= K_UP and min(levels[-K_UP:]) > prev:
        new = min(levels[-K_UP:])
    elif len(levels) >= K_DOWN and max(levels[-K_DOWN:]) < prev:
        new = max(levels[-K_DOWN:])
    else:
        new = prev
    st.session_state[key] = new
    return new, cur


def node_coords(nid):
    mc = st.session_state.setdefault("meta_cache", {})
    if nid not in mc or time.time() - mc[nid][0] > 300:
        meta = fb_get(f"nodes/{nid}/meta") or {}
        lat, lon = meta.get("lat"), meta.get("lon")
        if lat is None or lon is None:
            rng = np.random.default_rng(zlib.crc32(nid.encode()) & 0xFFFFFFFF)
            lat, lon = BASE_LAT + rng.uniform(-0.02, 0.02), BASE_LON + rng.uniform(-0.02, 0.02)
        mc[nid] = (time.time(), float(lat), float(lon))
    return mc[nid][1], mc[nid][2]


def fleet_row(nid, sensors, network, live_flag):
    sc = float(threat_score(sensors["soil_moisture_pct"], sensors["rainfall_pct"], sensors["tilt_angle_deg"], sensors["vibration_pulses"]))
    fos = float(factor_of_safety(sensors["tilt_angle_deg"], sensors["soil_moisture_pct"], SOIL))
    t = classify(sc, sensors["tilt_angle_deg"], sensors["soil_moisture_pct"], sensors["rainfall_pct"], fos, TH)
    lat, lon = node_coords(nid)
    return {"node": nid, "tier": t["name"], "level": t["level"], "threat": round(sc, 1), "fos": round(fos, 2),
            "moisture_pct": round(sensors["soil_moisture_pct"], 1), "rain_pct": round(sensors["rainfall_pct"], 1),
            "tilt_deg": round(sensors["tilt_angle_deg"], 1), "uplink": network.get("active_interface", "-"),
            "rssi": network.get("rssi"), "live": live_flag, "lat": lat, "lon": lon}


def get_fleet(sel_sensors, sel_network, sel_live):
    s = st.session_state
    ids = list_nodes() or [NODE_ID]
    demo = not sel_live and len(ids) <= 1
    if demo:
        ids = [NODE_ID] + [f"demo-node-{i:02d}" for i in (2, 3, 4)]
    cache = s.get("fleet_cache")
    if not cache or cache["key"] != (tuple(ids), NODE_ID, demo) or time.time() - cache["ts"] > 20:
        rows = []
        for nid in ids:
            if nid == NODE_ID:
                continue
            if demo:
                rng = np.random.default_rng(zlib.crc32(nid.encode()) & 0xFFFFFFFF)
                sn = dict(sel_sensors)
                sn["soil_moisture_pct"] = float(np.clip(sel_sensors["soil_moisture_pct"] + rng.normal(0, 15), 5, 98))
                sn["rainfall_pct"] = float(np.clip(sel_sensors["rainfall_pct"] + rng.normal(0, 15), 0, 100))
                sn["tilt_angle_deg"] = float(np.clip(sel_sensors["tilt_angle_deg"] + rng.normal(0, 6), 0.5, 40))
                sn["vibration_pulses"] = int(rng.integers(0, 6))
                rows.append(fleet_row(nid, sn, {"active_interface": "DEMO", "rssi": int(rng.integers(-90, -55))}, False))
            else:
                sn, net, lv, _ = parse_live(fb_get(f"nodes/{nid}/live"))
                rows.append(fleet_row(nid, sn, net, lv))
        s["fleet_cache"] = {"key": (tuple(ids), NODE_ID, demo), "ts": time.time(), "rows": rows}
    rows = [fleet_row(NODE_ID, sel_sensors, sel_network, sel_live)] + s["fleet_cache"]["rows"]
    return pd.DataFrame(rows).sort_values(["level", "threat"], ascending=False).reset_index(drop=True), demo


def node_health(df_raw, hist_is_real, sensors, extras, network, is_live):
    checks, score = [], 100

    def add(name, status, detail, penalty=0):
        nonlocal score
        checks.append({"check": name, "status": status, "detail": detail})
        score -= penalty

    if not is_live:
        add("Live connection", "bad", "No live data - showing demo values", 30)
    else:
        add("Live connection", "ok", "Receiving live telemetry")
    age = interval = None
    gaps = 0
    if hist_is_real:
        age = (pd.Timestamp.now(tz="UTC") - df_raw["timestamp"].iloc[-1]).total_seconds()
        dts = df_raw["timestamp"].diff().dt.total_seconds().dropna()
        interval = float(dts.median()) if len(dts) else None
        gaps = int((dts > 3 * interval).sum()) if interval and interval > 0 else 0
        if age > 600:
            add("Data freshness", "bad", f"Last record {age / 60:.0f} min ago", 40)
        elif age > 120:
            add("Data freshness", "warn", f"Last record {age / 60:.1f} min ago", 20)
        else:
            add("Data freshness", "ok", f"Last record {age:.0f} s ago")
        add("Sampling gaps", "ok" if gaps == 0 else "warn",
            f"{gaps} gap(s) > 3x median interval ({interval:.0f} s)" if interval else "n/a", min(5 * gaps, 20))
    else:
        add("Data freshness", "warn", "No device history - synthetic demo history in use", 10)

    tail = df_raw.tail(20)
    flat = [c for c in ("soil_moisture_pct", "temperature_c", "humidity_pct", "tilt_angle_deg")
            if len(tail) >= 20 and tail[c].std() < 1e-9]
    add("Stuck sensors", "ok" if not flat else "warn", "No flat-lined channels" if not flat else "Flat-lined: " + ", ".join(flat), 10 * len(flat))
    rules = {"soil_moisture_pct": (0, 100), "rainfall_pct": (0, 100), "humidity_pct": (0, 100),
             "tilt_angle_deg": (0, 90), "temperature_c": (-20, 70)}
    bad = [c for c, (lo, hi) in rules.items() if not lo <= sensors[c] <= hi]
    add("Range validation", "ok" if not bad else "bad", "All readings in physical range" if not bad else "Out of range: " + ", ".join(bad), 10 * len(bad))
    rssi = network.get("rssi")
    if isinstance(rssi, (int, float)) and rssi < 0:
        if rssi >= -67:
            add("Signal strength", "ok", f"{rssi} dBm (good)")
        elif rssi >= -80:
            add("Signal strength", "warn", f"{rssi} dBm (fair)", 5)
        else:
            add("Signal strength", "bad", f"{rssi} dBm (poor)", 15)
    for k, label in (("battery_v", "Battery"), ("battery_voltage", "Battery"), ("battery_pct", "Battery")):
        if k in extras:
            v = extras[k]
            low = v < 3.5 if "v" in k else v < 20
            add(label, "bad" if low else "ok", f"{v:g}{' V' if 'v' in k else ' %'}", 15 if low else 0)
            break
    for k in ("solar_v", "solar_voltage"):
        if k in extras:
            add("Solar input", "ok", f"{extras[k]:g} V")
            break
    score = max(score, 0)
    return {"score": score, "label": "Good" if score >= 85 else ("Degraded" if score >= 60 else "Poor"),
            "checks": checks, "age_s": age, "interval_s": interval, "gaps": gaps}


# ==============================================================================
# 6. FIGURES
# ==============================================================================
def gauge(value, title, vmax, steps, suffix="", fmt=".1f"):
    fig = go.Figure(go.Indicator(mode="gauge+number", value=value, title={"text": title, "font": {"size": 14}},
                                 number={"suffix": suffix, "valueformat": fmt},
                                 gauge={"axis": {"range": [0, vmax]}, "bar": {"color": "#F8FAFC", "thickness": 0.25},
                                        "steps": [{"range": [a, b], "color": c} for a, b, c in steps]}))
    fig.update_layout(height=230, margin=dict(l=20, r=20, t=50, b=10), **DARK)
    return fig


def radar(s):
    cats = ["Soil saturation", "Rainfall", "Tilt", "Tremor", "Humidity"]
    vals = [s["soil_moisture_pct"], s["rainfall_pct"], min(s["tilt_angle_deg"] / 45 * 100, 100),
            min(s["vibration_pulses"] * 10, 100), s["humidity_pct"]]
    fig = go.Figure(go.Scatterpolar(r=vals + vals[:1], theta=cats + cats[:1], fill="toself", line_color="#38BDF8"))
    fig.update_layout(polar=dict(radialaxis=dict(range=[0, 100], gridcolor="#334155"), bgcolor="#0F172A"),
                      showlegend=False, height=230, margin=dict(l=45, r=45, t=25, b=15), **DARK)
    return fig


def driver_waterfall(drivers, score):
    fig = go.Figure(go.Waterfall(
        x=[DRIVER_LABEL[k].title() for k in drivers] + ["Threat score"], y=list(drivers.values()) + [0],
        measure=["relative"] * len(drivers) + ["total"], connector={"line": {"color": "#475569"}},
        increasing={"marker": {"color": "#F59E0B"}}, totals={"marker": {"color": "#EF4444"}},
        text=[f"{v:.1f}" for v in drivers.values()] + [f"{score:.1f}"], textposition="outside"))
    fig.update_layout(title="What is driving the threat score", height=300, margin=dict(l=10, r=10, t=45, b=10),
                      yaxis=dict(range=[0, max(100, score * 1.15)]), showlegend=False, **DARK)
    return fig


def make_3d(s, mode, color):
    pitch, roll = np.radians(s["pitch_deg"]), np.radians(s["roll_deg"])
    g = np.linspace(-6, 6, 41)
    X, Y = np.meshgrid(g, g)
    Z = np.sin(X * 0.3 + pitch) * np.cos(Y * 0.3 + roll) * 2.2 + (X * np.sin(pitch) + Y * np.sin(roll)) * 1.5
    gy, gx = np.gradient(Z)
    grad = np.hypot(gx, gy)
    if mode == "Saturation stress":
        surf, cs, title = grad / max(grad.max(), 1e-9) * s["soil_moisture_pct"], "YlOrRd", "Stress index"
    else:
        surf, cs, title = Z, "Earth", "Elevation (m)"
    fig = go.Figure(go.Surface(x=X, y=Y, z=Z, surfacecolor=surf, colorscale=cs, colorbar=dict(title=title, len=0.6)))
    zc = float(Z[20, 20])
    fig.add_trace(go.Scatter3d(x=[0, np.sin(roll) * 4], y=[0, np.sin(pitch) * 4], z=[zc, zc + 3], mode="lines+markers",
                               line=dict(color="#EF4444", width=7), marker=dict(size=5, color="#F8FAFC"), name="Shear vector"))
    fig.add_trace(go.Scatter3d(x=[0], y=[0], z=[zc + 0.2], mode="markers+text", text=["Node"],
                               marker=dict(size=9, color=color, symbol="diamond"), name="Sensor node"))
    fig.update_layout(scene=dict(xaxis=dict(title="East-West (m)", backgroundcolor="#0F172A", gridcolor="#334155"),
                                 yaxis=dict(title="North-South (m)", backgroundcolor="#0F172A", gridcolor="#334155"),
                                 zaxis=dict(title="Elevation (m)", backgroundcolor="#0F172A", gridcolor="#334155"),
                                 camera=dict(eye=dict(x=1.6, y=1.6, z=1.2))),
                      height=520, margin=dict(l=0, r=0, t=10, b=0), paper_bgcolor="rgba(0,0,0,0)", font_color="#F1F5F9")
    return fig


def trend_figure(df, th, smooth=0, show_anom=True, fc=None, marker_ts=None):
    d = df.copy()
    if smooth and smooth > 1:
        for c in ("soil_moisture_pct", "rainfall_pct", "tilt_angle_deg", "threat_score"):
            d[c] = d[c].ewm(span=smooth, adjust=False).mean()
    fig = make_subplots(rows=3, cols=1, shared_xaxes=True, vertical_spacing=0.06,
                        subplot_titles=("Soil saturation & rainfall (%)", "Slope tilt (°)", "Tremor pulses & threat score"),
                        specs=[[{}], [{}], [{"secondary_y": True}]])
    t = d["timestamp"]
    fig.add_trace(go.Scatter(x=t, y=d["soil_moisture_pct"], name="Soil moisture", line=dict(color="#38BDF8")), 1, 1)
    fig.add_trace(go.Scatter(x=t, y=d["rainfall_pct"], name="Rainfall", line=dict(color="#818CF8")), 1, 1)
    fig.add_hline(y=80, line_dash="dash", line_color="#EF4444", row=1, col=1)
    fig.add_trace(go.Scatter(x=t, y=d["tilt_angle_deg"], name="Tilt", line=dict(color="#F59E0B")), 2, 1)
    fig.add_hline(y=th["tilt_high"], line_dash="dot", line_color="#F97316", row=2, col=1)
    fig.add_hline(y=th["tilt_crit"], line_dash="dash", line_color="#EF4444", row=2, col=1)
    if show_anom:
        m = flag_anomalies(df["tilt_angle_deg"])
        if m.any():
            fig.add_trace(go.Scatter(x=t[m], y=d["tilt_angle_deg"][m], mode="markers", name="Tilt anomaly",
                                     marker=dict(color="#F43F5E", size=11, symbol="x")), 2, 1)
    if fc and fc["slope"] > 0:
        t0 = t.iloc[-1]
        fig.add_trace(go.Scatter(x=[t0, t0 + pd.Timedelta(minutes=30)], y=[fc["cur"], fc["cur"] + fc["slope"] * 30], mode="lines",
                                 name="30-min trend", line=dict(color="#FACC15", dash="dash")), 2, 1)
    fig.add_trace(go.Bar(x=t, y=d["vibration_pulses"], name="Tremor", marker_color="#F43F5E", opacity=0.55), 3, 1)
    fig.add_trace(go.Scatter(x=t, y=d["threat_score"], name="Threat score", line=dict(color="#10B981", width=3)), 3, 1, secondary_y=True)
    if marker_ts is not None:
        fig.add_shape(type="line", x0=marker_ts, x1=marker_ts, y0=0, y1=1, yref="paper", line=dict(color="#F8FAFC", width=2, dash="dot"))
    fig.update_layout(height=720, margin=dict(l=10, r=10, t=40, b=10), legend=dict(orientation="h", y=-0.06), **DARK)
    return fig


def threat_timeline(df):
    fig = go.Figure()
    for a, b, c in ((0, 35, "#10B981"), (35, 55, "#F59E0B"), (55, 75, "#F97316"), (75, 100, "#EF4444")):
        fig.add_hrect(y0=a, y1=b, fillcolor=c, opacity=0.12, line_width=0)
    fig.add_trace(go.Scatter(x=df["timestamp"], y=df["threat_score"], line=dict(color="#F8FAFC", width=2), name="Threat score"))
    fig.update_layout(title="Threat score vs risk bands", yaxis=dict(range=[0, 100]), height=300, margin=dict(l=10, r=10, t=40, b=10), **DARK)
    return fig


def fan_chart(df, col, mc, limit, title):
    fig = go.Figure()
    d = df.tail(60)
    fig.add_trace(go.Scatter(x=d["timestamp"], y=d[col], name="Observed", line=dict(color="#38BDF8", width=2)))
    if mc:
        t0 = df["timestamp"].iloc[-1]
        xs = [t0 + pd.Timedelta(minutes=float(m)) for m in mc["t"]]
        fig.add_trace(go.Scatter(x=xs + xs[::-1], y=list(mc["p90"]) + list(mc["p10"][::-1]), fill="toself",
                                 fillcolor="rgba(250,204,21,0.18)", line=dict(width=0), name="P10-P90 band", hoverinfo="skip"))
        fig.add_trace(go.Scatter(x=xs, y=mc["p50"], name="Median", line=dict(color="#FACC15", dash="dash")))
    fig.add_hline(y=limit, line_dash="dash", line_color="#EF4444", annotation_text="limit")
    fig.update_layout(title=title, height=340, margin=dict(l=10, r=10, t=45, b=10), legend=dict(orientation="h", y=-0.15), **DARK)
    return fig


def fos_heatmap(soil, live_pt, sim_pt):
    moist, tilt = np.linspace(0, 100, 41), np.linspace(2, 45, 44)
    M, T = np.meshgrid(moist, tilt)
    F = factor_of_safety(T, M, soil)
    fig = go.Figure()
    fig.add_trace(go.Heatmap(x=moist, y=tilt, z=F, colorscale="RdYlGn", zmin=0, zmax=3, colorbar=dict(title="FoS")))
    fig.add_trace(go.Contour(x=moist, y=tilt, z=F, showscale=False, contours=dict(start=1, end=1, size=1, coloring="lines", showlabels=True),
                             line=dict(color="white", width=3), name="FoS = 1"))
    fig.add_trace(go.Scatter(x=[live_pt[0]], y=[live_pt[1]], mode="markers", name="Live",
                             marker=dict(size=14, color="#38BDF8", line=dict(color="white", width=2))))
    fig.add_trace(go.Scatter(x=[sim_pt[0]], y=[sim_pt[1]], mode="markers", name="Scenario",
                             marker=dict(size=14, color="#F8FAFC", symbol="star", line=dict(color="black", width=1))))
    fig.update_layout(title="Stability map (white line = FoS 1)", xaxis_title="Soil saturation (%)", yaxis_title="Slope tilt (°)",
                      height=440, margin=dict(l=10, r=10, t=45, b=10), **DARK)
    return fig


def tornado(s, soil):
    base = float(factor_of_safety(s["tilt_angle_deg"], s["soil_moisture_pct"], soil))
    items = []
    for label, kind, key in (("Cohesion c'", "soil", "c"), ("Friction angle", "soil", "phi"), ("Unit weight", "soil", "gamma"),
                             ("Failure depth", "soil", "depth"), ("Soil saturation", "sens", "m"), ("Slope tilt", "sens", "t")):
        deltas = []
        for mult in (0.8, 1.2):
            s2, tilt, m = dict(soil), s["tilt_angle_deg"], s["soil_moisture_pct"]
            if kind == "soil":
                s2[key] = soil[key] * mult
            elif key == "m":
                m = min(m * mult, 100)
            else:
                tilt = tilt * mult
            deltas.append(float(factor_of_safety(tilt, m, s2)) - base)
        items.append((label, *deltas))
    items.sort(key=lambda x: abs(x[1] - x[2]))
    fig = go.Figure()
    fig.add_trace(go.Bar(y=[i[0] for i in items], x=[i[1] for i in items], orientation="h", name="-20 %", marker_color="#38BDF8"))
    fig.add_trace(go.Bar(y=[i[0] for i in items], x=[i[2] for i in items], orientation="h", name="+20 %", marker_color="#F97316"))
    fig.update_layout(barmode="overlay", title=f"FoS sensitivity to ±20 % change (base FoS {base:.2f})", xaxis_title="Change in FoS",
                      height=330, margin=dict(l=10, r=10, t=45, b=10), **DARK)
    return fig


def rel_hist(rel):
    fig = go.Figure(go.Histogram(x=rel["f"], nbinsx=50, marker_color="#38BDF8"))
    fig.add_vline(x=1.0, line_dash="dash", line_color="#EF4444", annotation_text="FoS = 1")
    fig.update_layout(title=f"Monte-Carlo FoS distribution - P(failure) = {rel['pf'] * 100:.1f} %", xaxis_title="Factor of Safety",
                      height=330, margin=dict(l=10, r=10, t=45, b=10), **DARK)
    return fig


def fleet_map(fl):
    Trace = go.Scattermap if hasattr(go, "Scattermap") else go.Scattermapbox
    fig = go.Figure(Trace(lat=fl["lat"], lon=fl["lon"], mode="markers+text", text=fl["node"], textposition="top right",
                          marker=dict(size=20, color=[TIERS[t]["color"] for t in fl["tier"]]),
                          hovertext=[f"{r.node}<br>{r.tier} - threat {r.threat}<br>FoS {r.fos}" for r in fl.itertuples()], hoverinfo="text"))
    layout = dict(style="carto-darkmatter", center=dict(lat=float(fl["lat"].mean()), lon=float(fl["lon"].mean())), zoom=11)
    fig.update_layout(**({"map": layout} if hasattr(go, "Scattermap") else {"mapbox": layout}),
                      height=430, margin=dict(l=0, r=0, t=0, b=0), paper_bgcolor="rgba(0,0,0,0)")
    return fig


# ==============================================================================
# 7. CONTEXT BUILDER + NARRATIVE
# ==============================================================================
def make_narrative(S):
    t, s = S.tier, S.sensors
    top = max(S.drivers, key=S.drivers.get)
    parts = [f"{NODE_ID} is at {t['name']} status with a threat score of {S.score:.0f}/100 and a factor of safety of {S.fos:.2f}."]
    if S.pending:
        parts.append(f"The raw signal currently reads {S.raw_tier['name']}, held back by alert debouncing.")
    parts.append(f"The largest contributor is {DRIVER_LABEL[top]} ({S.drivers[top]:.0f} of {S.score:.0f} points).")
    if S.mc_tilt and S.mc_tilt["slope"] * 60 > 0.05:
        parts.append(f"Tilt is rising at about {S.mc_tilt['slope'] * 60:.2f} deg/h; at this rate the critical limit of "
                     f"{TH['tilt_crit']:.0f} deg would be reached in {fmt_eta(S.fc_tilt)}.")
    elif S.mc_tilt:
        parts.append("Tilt is essentially stable.")
    if S.mc_moist and S.mc_moist["slope"] * 60 > 0.5:
        parts.append(f"Soil saturation is climbing about {S.mc_moist['slope'] * 60:.1f} %/h (critical level in {fmt_eta(S.fc_moist)}).")
    parts.append(f"Reliability analysis puts the probability of failure (FoS below 1) at {S.rel['pf'] * 100:.1f}% under current conditions"
                 + (f", and {S.p_fos * 100:.0f}% within the next {FC_H} minutes if trends continue." if S.p_fos is not None else "."))
    parts.append(f"Data quality is {S.health['label'].lower()} ({S.health['score']:.0f}/100).")
    parts.append("Priority action: " + PROTOCOL[t["name"]][0])
    return " ".join(parts)


def build_context():
    S = SimpleNamespace()
    live = fb_get(f"nodes/{NODE_ID}/live")
    S.sensors, S.network, S.is_live, S.extras = parse_live(live)
    raw = fb_get(f"nodes/{NODE_ID}/history", {"orderBy": '"$key"', "limitToLast": 400})
    if raw is None:
        raw = fb_get(f"nodes/{NODE_ID}/history")
    df_raw = normalize_history(raw, S.sensors)
    S.hist_is_real = df_raw is not None
    if df_raw is None:
        df_raw = synth_history(S.sensors, NODE_ID)
    S.df_raw = df_raw
    df = df_raw
    if S.is_live:
        now = pd.Timestamp.now(tz="UTC")
        if now > df["timestamp"].iloc[-1]:
            df = pd.concat([df, pd.DataFrame([{"timestamp": now, **S.sensors}])], ignore_index=True)
    S.df = enrich(df, TH, SOIL)

    s = S.sensors
    S.score = float(threat_score(s["soil_moisture_pct"], s["rainfall_pct"], s["tilt_angle_deg"], s["vibration_pulses"]))
    S.fos = float(factor_of_safety(s["tilt_angle_deg"], s["soil_moisture_pct"], SOIL))
    S.raw_tier = classify(S.score, s["tilt_angle_deg"], s["soil_moisture_pct"], s["rainfall_pct"], S.fos, TH)
    lvl, raw_lvl = apply_debounce([TIERS[n]["level"] for n in S.df["risk_tier"]])
    S.pending = lvl != raw_lvl
    S.tier = {"name": LEVEL_NAME[lvl], **TIERS[LEVEL_NAME[lvl]], "score": S.score, "fos": S.fos}
    S.drivers = score_drivers(s)

    S.fc_tilt = forecast_eta(S.df, "tilt_angle_deg", TH["tilt_crit"])
    S.fc_moist = forecast_eta(S.df, "soil_moisture_pct", TH["moist_crit"])
    S.mc_tilt = mc_forecast(S.df, "tilt_angle_deg", FC_H)
    S.mc_moist = mc_forecast(S.df, "soil_moisture_pct", FC_H)
    S.p_tilt, S.p_moist = p_breach(S.mc_tilt, TH["tilt_crit"]), p_breach(S.mc_moist, TH["moist_crit"])
    S.p_fos = None
    if S.mc_tilt and S.mc_moist and S.mc_tilt["paths"].shape == S.mc_moist["paths"].shape:
        f = factor_of_safety(np.clip(S.mc_tilt["paths"], 0, 60), np.clip(S.mc_moist["paths"], 0, 100), SOIL)
        S.p_fos = float((f.min(axis=1) < 1).mean())
    S.rel = reliability(s["tilt_angle_deg"], s["soil_moisture_pct"], SOIL, CV)
    S.health = node_health(df_raw, S.hist_is_real, s, S.extras, S.network, S.is_live)
    S.fleet, S.fleet_demo = get_fleet(s, S.network, S.is_live)
    S.source = "live Firebase stream" if S.is_live else "demo data"
    S.narrative = make_narrative(S)
    S.stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    return S


def log_and_notify(S):
    cur = S.tier["name"]
    prev = st.session_state.last_tier
    if prev is None:
        log_event("INFO", f"Monitoring session started - status {cur}")
    elif prev != cur:
        up = S.tier["level"] > TIERS[prev]["level"]
        log_event("ALERT" if up else "INFO", f"Risk {'escalated' if up else 'eased'}: {prev} to {cur} (score {S.score:.1f}, FoS {S.fos:.2f})")
    st.session_state.last_tier = cur
    if flag_anomalies(S.df["tilt_angle_deg"]).iloc[-1] and st.session_state.last_anom_ts != S.df["timestamp"].iloc[-1]:
        st.session_state.last_anom_ts = S.df["timestamp"].iloc[-1]
        log_event("ANOMALY", f"Abrupt tilt change detected ({S.sensors['tilt_angle_deg']:.1f} deg)")

    last = st.session_state.setdefault("last_notify", {"level": 0, "ts": 0.0})
    if S.tier["level"] < last["level"]:
        last["level"] = S.tier["level"]
    if NOTIFY_ON and WEBHOOK and S.tier["level"] >= NOTIFY_MIN and (
            S.tier["level"] > last["level"] or time.time() - last["ts"] >= NOTIFY_COOLDOWN * 60):
        msg = f"GeoAlert {cur} at {NODE_ID}: threat {S.score:.0f}/100, FoS {S.fos:.2f}. " + PROTOCOL[cur][0]
        try:
            requests.post(WEBHOOK, json={"text": msg, "content": msg, "node": NODE_ID, "tier": cur, "score": round(S.score, 1),
                                         "fos": round(S.fos, 2)}, timeout=6).raise_for_status()
            log_event("NOTIFY", "Webhook notification sent")
        except Exception as err:
            log_event("ERROR", f"Webhook failed: {err}")
        last["level"], last["ts"] = S.tier["level"], time.time()


# ==============================================================================
# 8. EXPORT BUILDERS
# ==============================================================================
def snapshot(S):
    return {
        "node": NODE_ID, "generated_utc": S.stamp, "source": S.source, "uplink": S.network, "sensors": S.sensors,
        "risk": {"tier": S.tier["name"], "raw_tier": S.raw_tier["name"], "score": round(S.score, 2), "fos": round(S.fos, 3),
                 "probability_of_failure": round(S.rel["pf"], 4), "reliability_index": None if not np.isfinite(S.rel["beta"]) else round(S.rel["beta"], 2)},
        "forecast": {"horizon_min": FC_H, "tilt_eta": fmt_eta(S.fc_tilt), "moisture_eta": fmt_eta(S.fc_moist),
                     "p_tilt_critical": S.p_tilt, "p_saturation_critical": S.p_moist, "p_fos_below_1": S.p_fos},
        "health": {"score": S.health["score"], "label": S.health["label"], "checks": S.health["checks"]},
        "thresholds": TH, "soil_model": SOIL, "narrative": S.narrative,
    }


def forecast_frame(S):
    if not S.mc_tilt:
        return pd.DataFrame()
    t0 = S.df["timestamp"].iloc[-1].tz_convert(None)
    out = pd.DataFrame({"minutes_ahead": S.mc_tilt["t"], "timestamp_utc": [t0 + pd.Timedelta(minutes=float(m)) for m in S.mc_tilt["t"]],
                        "tilt_p10": S.mc_tilt["p10"], "tilt_p50": S.mc_tilt["p50"], "tilt_p90": S.mc_tilt["p90"]})
    if S.mc_moist:
        out["moisture_p10"], out["moisture_p50"], out["moisture_p90"] = S.mc_moist["p10"], S.mc_moist["p50"], S.mc_moist["p90"]
    return out


def kpi_rows(S):
    return [("Node", NODE_ID), ("Risk tier", S.tier["name"]), ("Threat score (0-100)", round(S.score, 1)),
            ("Factor of safety", round(S.fos, 2)), ("Probability of failure now", f"{S.rel['pf'] * 100:.1f} %"),
            (f"P(FoS<1) within {FC_H} min", "n/a" if S.p_fos is None else f"{S.p_fos * 100:.0f} %"),
            ("Tilt to critical ETA", fmt_eta(S.fc_tilt)), ("Saturation to critical ETA", fmt_eta(S.fc_moist)),
            ("Data health", f"{S.health['label']} ({S.health['score']:.0f}/100)"), ("Data source", S.source), ("Generated (UTC)", S.stamp)]


def build_xlsx(S):
    from openpyxl.chart import LineChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ex = S.df[EXPORT_COLS].copy()
    ex["timestamp"] = ex["timestamp"].dt.tz_convert(None)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame(kpi_rows(S), columns=["Metric", "Value"]).to_excel(w, sheet_name="Executive_Summary", index=False, startrow=5)
        ex.to_excel(w, sheet_name="Time_Series", index=False)
        fc = forecast_frame(S)
        (fc if len(fc) else pd.DataFrame({"note": ["Not enough history for a forecast"]})).to_excel(w, sheet_name="Forecast", index=False)
        S.fleet.drop(columns=["level"]).to_excel(w, sheet_name="Fleet", index=False)
        pd.DataFrame(S.health["checks"]).to_excel(w, sheet_name="Node_Health", index=False)
        al = pd.DataFrame(st.session_state.alert_log) if st.session_state.alert_log else pd.DataFrame(columns=["time_utc", "level", "event", "ack"])
        al.to_excel(w, sheet_name="Alert_Log", index=False)

        head = PatternFill("solid", fgColor="1E293B")
        for ws in w.book.worksheets:
            hdr_row = 6 if ws.title == "Executive_Summary" else 1
            for c in ws[hdr_row]:
                c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head, Alignment(horizontal="center")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:120])
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, width + 2), 48)
            if ws.title != "Executive_Summary":
                ws.freeze_panes = "A2"
        es = w.book["Executive_Summary"]
        es["A1"], es["A1"].font = "GeoAlert Incident Summary", Font(bold=True, size=16)
        es["A2"] = f"{NODE_ID} | {S.stamp} UTC | {S.source}"
        es["A3"], es["A3"].alignment = S.narrative, Alignment(wrap_text=True, vertical="top")
        es.merge_cells("A3:D5")
        es.column_dimensions["B"].width = 34
        tier_cell = next(r[1] for r in es.iter_rows(min_row=7, max_col=2) if r[0].value == "Risk tier")
        tier_cell.fill = PatternFill("solid", fgColor=S.tier["color"].lstrip("#"))
        tier_cell.font = Font(bold=True, color="FFFFFF")

        ts = w.book["Time_Series"]
        cols = list(ex.columns)
        tcol, ccol = cols.index("risk_tier") + 1, cols.index("threat_score") + 1
        for row in ts.iter_rows(min_row=2, min_col=tcol, max_col=tcol):
            row[0].fill = PatternFill("solid", fgColor=TIERS.get(row[0].value, TIERS["NORMAL"])["color"].lstrip("#"))
        L = get_column_letter(ccol)
        ts.conditional_formatting.add(f"{L}2:{L}{ts.max_row}", ColorScaleRule(
            start_type="num", start_value=0, start_color="10B981", mid_type="num", mid_value=50, mid_color="F59E0B",
            end_type="num", end_value=100, end_color="EF4444"))
        for cell in ts["A"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        if len(ex) > 2:
            for title, names, anchor in (("Moisture, rainfall & tilt", ("soil_moisture_pct", "rainfall_pct", "tilt_angle_deg"), "B2"),
                                         ("Threat score & factor of safety", ("threat_score", "fos"), "B22")):
                ch = LineChart()
                ch.title, ch.height, ch.width = title, 9, 24
                for n in names:
                    ch.add_data(Reference(ts, min_col=cols.index(n) + 1, min_row=1, max_row=ts.max_row), titles_from_data=True)
                ch.set_categories(Reference(ts, min_col=1, min_row=2, max_row=ts.max_row))
                cs = w.book["Charts"] if "Charts" in w.book.sheetnames else w.book.create_sheet("Charts")
                cs.add_chart(ch, anchor)
    return buf.getvalue()


def build_geojson(S):
    feats = [{"type": "Feature", "geometry": {"type": "Point", "coordinates": [r.lon, r.lat]},
              "properties": {"node": r.node, "tier": r.tier, "threat": r.threat, "fos": r.fos, "moisture_pct": r.moisture_pct,
                             "tilt_deg": r.tilt_deg, "live": bool(r.live), "generated_utc": S.stamp}} for r in S.fleet.itertuples()]
    return json.dumps({"type": "FeatureCollection", "features": feats}, indent=2).encode()


def build_html(S, figs):
    c = S.tier["color"]
    cards = "".join(f"<div class='k'><span>{k}</span><b>{v}</b></div>" for k, v in kpi_rows(S)[2:9])
    plots = "".join(f.to_html(full_html=False, include_plotlyjs=True if i == 0 else False) for i, f in enumerate(figs))
    fleet = "".join(f"<tr><td>{r.node}</td><td style='color:{TIERS[r.tier]['color']}'><b>{r.tier}</b></td><td>{r.threat}</td>"
                    f"<td>{r.fos}</td><td>{r.moisture_pct}</td><td>{r.tilt_deg}</td></tr>" for r in S.fleet.itertuples())
    checks = "".join(f"<tr><td>{x['check']}</td><td>{x['status'].upper()}</td><td>{x['detail']}</td></tr>" for x in S.health["checks"])
    acts = "".join(f"<li>{a}</li>" for a in PROTOCOL[S.tier["name"]])
    log = "".join(f"<tr><td>{a['time_utc']}</td><td>{a['level']}</td><td>{a['event']}</td></tr>" for a in st.session_state.alert_log[:25])
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>GeoAlert Report - {NODE_ID}</title><style>
body{{background:#0B1120;color:#F1F5F9;font-family:Inter,Segoe UI,Arial,sans-serif;margin:0;padding:28px;max-width:1200px;margin:auto}}
h1{{margin:0}} h2{{border-bottom:1px solid #334155;padding-bottom:6px;margin-top:32px}}
.hero{{border-left:8px solid {c};background:#1E293B;padding:18px 22px;border-radius:12px;margin:16px 0}}
.tier{{color:{c};font-size:2.2rem;font-weight:800}} .grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin:14px 0}}
.k{{background:#1E293B;border:1px solid #334155;border-radius:10px;padding:12px}} .k span{{display:block;color:#94A3B8;font-size:.75rem;text-transform:uppercase}}
.k b{{font-size:1.3rem}} table{{border-collapse:collapse;width:100%;margin:10px 0}} td,th{{border:1px solid #334155;padding:6px 10px;text-align:left;font-size:.9rem}}
th{{background:#1E293B}} .brief{{background:#111C33;border-left:4px solid #38BDF8;padding:12px 16px;border-radius:8px;line-height:1.55}}</style></head><body>
<h1>GeoAlert Incident Report</h1><p style="color:#94A3B8">{NODE_ID} · generated {S.stamp} UTC · {S.source}</p>
<div class="hero"><div class="tier">{S.tier['name']}</div><div class="grid">{cards}</div></div>
<h2>Situation summary</h2><div class="brief">{S.narrative}</div>
<h2>Recommended actions</h2><ul>{acts}</ul>
<h2>Charts</h2>{plots}
<h2>Fleet status</h2><table><tr><th>Node</th><th>Tier</th><th>Threat</th><th>FoS</th><th>Saturation %</th><th>Tilt °</th></tr>{fleet}</table>
<h2>Node health</h2><table><tr><th>Check</th><th>Status</th><th>Detail</th></tr>{checks}</table>
<h2>Alert log</h2><table><tr><th>Time (UTC)</th><th>Level</th><th>Event</th></tr>{log}</table>
<p style="color:#64748B;font-size:.8rem;margin-top:30px">Decision-support output. FoS uses an infinite-slope model with the soil parameters configured in the dashboard; forecasts are linear-trend Monte-Carlo estimates. Validate against site investigation before operational decisions.</p>
</body></html>""".encode()


def _pdf_txt(s):
    return str(s).encode("cp1252", "replace").decode("cp1252")


def build_pdf(S):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    df = S.df.tail(120)
    t = df["timestamp"].dt.tz_convert(None)

    def to_img(fig, w_mm):
        b = io.BytesIO()
        fig.savefig(b, format="png", dpi=150)
        wi, hi = fig.get_size_inches()
        plt.close(fig)
        b.seek(0)
        return Image(b, width=w_mm * mm, height=w_mm * mm * hi / wi)

    f1, ax = plt.subplots(3, 1, figsize=(8, 6.6), sharex=True)
    ax[0].plot(t, df["soil_moisture_pct"], label="Soil moisture", color="#0284C7")
    ax[0].plot(t, df["rainfall_pct"], label="Rainfall", color="#6366F1")
    ax[0].axhline(80, ls="--", color="red", lw=1)
    ax[0].set_ylabel("%"); ax[0].legend(loc="upper left", fontsize=8)
    ax[1].plot(t, df["tilt_angle_deg"], color="#D97706"); ax[1].axhline(TH["tilt_crit"], ls="--", color="red", lw=1)
    ax[1].set_ylabel("Tilt (deg)")
    ax[2].plot(t, df["threat_score"], color="#059669")
    for a, b, c in ((0, 35, "#10B981"), (35, 55, "#F59E0B"), (55, 75, "#F97316"), (75, 100, "#EF4444")):
        ax[2].axhspan(a, b, color=c, alpha=0.15)
    ax[2].set_ylabel("Threat score"); ax[2].set_ylim(0, 100)
    f1.autofmt_xdate(); f1.tight_layout()

    f2, ax2 = plt.subplots(1, 3, figsize=(10, 3.1))
    for a, (col, mc, lim, ttl) in zip(ax2[:2], [("tilt_angle_deg", S.mc_tilt, TH["tilt_crit"], "Tilt forecast (deg)"),
                                                 ("soil_moisture_pct", S.mc_moist, TH["moist_crit"], "Saturation forecast (%)")]):
        d = S.df.tail(60)
        tt = d["timestamp"].dt.tz_convert(None)
        a.plot(tt, d[col], color="#0284C7")
        if mc:
            xs = tt.iloc[-1] + pd.to_timedelta(mc["t"], unit="m")
            a.fill_between(xs, mc["p10"], mc["p90"], color="#FACC15", alpha=0.3)
            a.plot(xs, mc["p50"], "--", color="#CA8A04")
        a.axhline(lim, color="red", ls="--", lw=1); a.set_title(ttl, fontsize=9); a.tick_params(labelsize=6)
        for lbl in a.get_xticklabels():
            lbl.set_rotation(30)
    ax2[2].hist(S.rel["f"], bins=40, color="#0284C7"); ax2[2].axvline(1, color="red", ls="--")
    ax2[2].set_title(f"FoS distribution (Pf = {S.rel['pf'] * 100:.1f}%)", fontsize=9); ax2[2].tick_params(labelsize=7)
    f2.tight_layout()

    ss = getSampleStyleSheet()
    tier_col = colors.HexColor(S.tier["color"])
    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8),
                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)])
    story = [Paragraph("GeoAlert Incident Report", ss["Title"]),
             Paragraph(_pdf_txt(f"{NODE_ID} | {S.stamp} UTC | {S.source}"), ss["Normal"]), Spacer(1, 8)]
    head = Table([[f"RISK LEVEL: {S.tier['name']}", f"Threat {S.score:.1f}/100    FoS {S.fos:.2f}    Pf {S.rel['pf'] * 100:.1f}%"]],
                 colWidths=[70 * mm, 100 * mm])
    head.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), tier_col), ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                              ("FONTSIZE", (0, 0), (-1, -1), 12), ("PADDING", (0, 0), (-1, -1), 8)]))
    story += [head, Spacer(1, 10), Paragraph("Situation summary", ss["Heading2"]), Paragraph(_pdf_txt(S.narrative), ss["BodyText"]),
              Spacer(1, 6), Paragraph("Recommended actions", ss["Heading2"])]
    story += [Paragraph(_pdf_txt(f"&bull; {a}"), ss["Normal"]) for a in PROTOCOL[S.tier["name"]]]
    kp = Table([["Metric", "Value"]] + [[k, _pdf_txt(v)] for k, v in kpi_rows(S)], colWidths=[85 * mm, 85 * mm])
    kp.setStyle(grid)
    story += [Spacer(1, 8), Paragraph("Key indicators", ss["Heading2"]), kp, Spacer(1, 6), Paragraph("Latest telemetry", ss["Heading2"])]
    tel = Table([["Sensor", "Value"]] + [[k, f"{v:g}"] for k, v in S.sensors.items()], colWidths=[85 * mm, 85 * mm])
    tel.setStyle(grid)
    story += [tel, PageBreak(), Paragraph("Trends", ss["Heading2"]), to_img(f1, 165), Spacer(1, 6),
              Paragraph("Forecast & reliability", ss["Heading2"]), to_img(f2, 175), PageBreak(),
              Paragraph("Fleet status", ss["Heading2"])]
    fl = Table([["Node", "Tier", "Threat", "FoS", "Sat %", "Tilt"]] +
               [[r.node, r.tier, r.threat, r.fos, r.moisture_pct, r.tilt_deg] for r in S.fleet.itertuples()],
              colWidths=[60 * mm, 28 * mm, 20 * mm, 20 * mm, 20 * mm, 20 * mm])
    fl.setStyle(grid)
    story += [fl, Spacer(1, 6), Paragraph("Node health", ss["Heading2"])]
    hl = Table([["Check", "Status", "Detail"]] + [[c["check"], c["status"].upper(), _pdf_txt(c["detail"])[:70]] for c in S.health["checks"]],
               colWidths=[40 * mm, 20 * mm, 110 * mm])
    hl.setStyle(grid)
    story += [hl, Spacer(1, 6), Paragraph("Recent alert log", ss["Heading2"])]
    if st.session_state.alert_log:
        al = Table([["Time (UTC)", "Level", "Event"]] + [[a["time_utc"], a["level"], _pdf_txt(a["event"])[:75]] for a in st.session_state.alert_log[:12]],
                   colWidths=[38 * mm, 22 * mm, 110 * mm])
        al.setStyle(grid)
        story.append(al)
    else:
        story.append(Paragraph("No events logged this session.", ss["Normal"]))
    story += [Spacer(1, 10), Paragraph(
        "<i>Decision-support output. FoS uses an infinite-slope model with the configured soil parameters; forecasts are linear-trend "
        "Monte-Carlo estimates. Validate against site investigation before operational decisions.</i>", ss["Normal"])]

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(18 * mm, 8 * mm, _pdf_txt(f"GeoAlert | {NODE_ID} | {S.stamp} UTC"))
        canvas.drawRightString(A4[0] - 18 * mm, 8 * mm, f"Page {doc.page}")

    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm).build(
        story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def build_package(S):
    files = {}
    snap = snapshot(S)
    files["xlsx"] = build_xlsx(S)
    files["csv"] = S.df[EXPORT_COLS].to_csv(index=False).encode()
    files["json"] = json.dumps({**snap, "fleet": json.loads(S.fleet.to_json(orient="records")),
                                "history": json.loads(S.df[EXPORT_COLS].tail(200).to_json(orient="records", date_format="iso")),
                                "alerts": st.session_state.alert_log}, indent=2, default=str).encode()
    files["geojson"] = build_geojson(S)
    figs = [trend_figure(S.df.tail(120), TH, 0, True, S.fc_tilt), threat_timeline(S.df.tail(120)),
            fan_chart(S.df, "tilt_angle_deg", S.mc_tilt, TH["tilt_crit"], "Tilt forecast (deg)"),
            fan_chart(S.df, "soil_moisture_pct", S.mc_moist, TH["moist_crit"], "Soil saturation forecast (%)"),
            tornado(S.sensors, SOIL), rel_hist(S.rel)]
    files["html"] = build_html(S, figs)
    try:
        files["pdf"] = build_pdf(S)
    except ImportError:
        files["pdf"] = None
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    names = {"xlsx": "workbook.xlsx", "csv": "timeseries.csv", "json": "snapshot.json", "geojson": "nodes.geojson",
             "html": "interactive_report.html", "pdf": "incident_report.pdf"}
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in files.items():
            if v:
                z.writestr(f"GeoAlert_{NODE_ID}_{stamp}/{names[k]}", v)
        z.writestr(f"GeoAlert_{NODE_ID}_{stamp}/README.txt",
                   f"GeoAlert export for {NODE_ID}\nGenerated {S.stamp} UTC ({S.source})\n\n{S.narrative}\n")
    files["zip"] = zb.getvalue()
    return {"stamp": stamp, "files": files, "built_utc": S.stamp}


# ==============================================================================
# 9. DASHBOARD TABS
# ==============================================================================
def show_df(df, **kw):
    st.dataframe(df, hide_index=True, **WIDE, **kw)


def tab_situation(S):
    s = S.sensors
    g1, g2, g3, g4 = st.columns(4)
    g1.plotly_chart(gauge(S.score, "Threat score", 100, [(0, 35, "#065F46"), (35, 55, "#92400E"), (55, 75, "#9A3412"), (75, 100, "#991B1B")]), **WIDE, key="g_threat")
    g2.plotly_chart(gauge(S.fos, "Factor of safety", 3, [(0, 1, "#991B1B"), (1, 1.25, "#9A3412"), (1.25, 1.5, "#92400E"), (1.5, 3, "#065F46")], fmt=".2f"), **WIDE, key="g_fos")
    g3.plotly_chart(gauge(S.rel["pf"] * 100, "Probability of failure", 100, [(0, 5, "#065F46"), (5, 20, "#92400E"), (20, 50, "#9A3412"), (50, 100, "#991B1B")], suffix=" %"), **WIDE, key="g_pf")
    g4.plotly_chart(radar(s), **WIDE, key="g_radar")

    st.markdown(f"<div class='brief'><b>Situation report</b><br>{S.narrative}</div>", unsafe_allow_html=True)

    f = st.columns(4)
    f[0].metric("Tilt to critical", fmt_eta(S.fc_tilt), f"{S.fc_tilt['slope'] * 60:+.2f} °/h" if S.fc_tilt else None, delta_color="inverse")
    f[1].metric("Saturation to critical", fmt_eta(S.fc_moist), f"{S.fc_moist['slope'] * 60:+.1f} %/h" if S.fc_moist else None, delta_color="inverse")
    f[2].metric(f"P(FoS<1) in {FC_H} min", "n/a" if S.p_fos is None else f"{S.p_fos * 100:.0f} %")
    f[3].metric("Reliability index β", "∞" if not np.isfinite(S.rel["beta"]) else f"{S.rel['beta']:.2f}")

    l, r = st.columns([3, 2])
    l.plotly_chart(driver_waterfall(S.drivers, S.score), **WIDE, key="waterfall")
    with r:
        st.markdown("#### Recommended protocol")
        for a in PROTOCOL[S.tier["name"]]:
            st.checkbox(a, key=f"proto_{NODE_ID}_{S.tier['name']}_{a}")
        st.markdown("#### Latest events")
        if st.session_state.alert_log:
            show_df(pd.DataFrame(st.session_state.alert_log[:5])[["time_utc", "level", "event"]])
        else:
            st.caption("No events yet.")


def tab_fleet(S):
    fl = S.fleet
    if S.fleet_demo:
        st.info("Only one node found - showing simulated neighbours (demo). Add nodes under `nodes/<id>/live` (and optional `meta {lat, lon}`) to see your real fleet.")
    k = st.columns(5)
    k[0].metric("Nodes", len(fl))
    for i, name in enumerate(LEVEL_NAME.values()):
        k[i + 1].metric(name.title(), int((fl["tier"] == name).sum()))
    st.plotly_chart(fleet_map(fl), **WIDE, key="fleet_map")
    show_df(fl.drop(columns=["level", "lat", "lon"]), column_config={
        "threat": st.column_config.ProgressColumn("Threat", min_value=0, max_value=100, format="%.0f"),
        "fos": st.column_config.NumberColumn("FoS", format="%.2f")})


def tab_twin(S):
    s = S.sensors
    st.subheader("Geological slope mesh & sensor orientation")
    cv, cs = st.columns([3, 1])
    with cv:
        mode = st.radio("Surface colouring", ["Elevation", "Saturation stress"], horizontal=True, key="mode3d")
        st.plotly_chart(make_3d(s, mode, S.tier["color"]), **WIDE, key="fig3d")
    with cs:
        st.markdown("#### Kinematics")
        st.info(f"**Pitch:** {s['pitch_deg']:.1f}°")
        st.info(f"**Roll:** {s['roll_deg']:.1f}°")
        st.info(f"**Yaw:** {s['yaw_deg']:.1f}°")
        st.info(f"**Resultant tilt:** {s['tilt_angle_deg']:.1f}°")
        st.metric("Factor of Safety", f"{S.fos:.2f}", f"{S.fos - 1.0:+.2f} vs limit")
        if S.fos < 1.0:
            st.warning("FoS below 1.0 - slope at limit-equilibrium failure!")


def tab_analytics(S):
    df = S.df
    n = len(df)
    o1, o2, o3 = st.columns(3)
    smooth = o1.slider("Smoothing (EWMA span)", 0, 30, 0, key="smooth")
    window = o2.slider("Window (samples)", 20, n, min(n, 90), key="win") if n >= 30 else n
    show_anom = o3.checkbox("Mark tilt anomalies", True, key="anom")
    dv = df.tail(window)
    st.plotly_chart(trend_figure(dv, TH, smooth, show_anom, S.fc_tilt), **WIDE, key="trend")

    cols = ["soil_moisture_pct", "rainfall_pct", "tilt_angle_deg", "vibration_pulses", "temperature_c", "humidity_pct"]
    mv = mahalanobis_scores(dv, cols)
    lim = chi_threshold(len(cols))
    fig_mv = go.Figure(go.Scatter(x=dv["timestamp"], y=mv, line=dict(color="#A78BFA"), name="Anomaly score"))
    fig_mv.add_hline(y=lim, line_dash="dash", line_color="#EF4444", annotation_text="99 % limit")
    flagged = mv > lim
    if flagged.any():
        fig_mv.add_trace(go.Scatter(x=dv["timestamp"][flagged], y=mv[flagged], mode="markers", marker=dict(color="#F43F5E", size=9), name="Anomalous"))
    fig_mv.update_layout(title="Multivariate anomaly score (all sensors jointly)", height=300, margin=dict(l=10, r=10, t=45, b=10), **DARK)
    c1, c2 = st.columns(2)
    c1.plotly_chart(fig_mv, **WIDE, key="mv")
    fig_ari = make_subplots(specs=[[{"secondary_y": True}]])
    fig_ari.add_trace(go.Scatter(x=dv["timestamp"], y=dv["antecedent_rain"], name="Antecedent rainfall (EWMA)", line=dict(color="#818CF8")), secondary_y=False)
    fig_ari.add_trace(go.Scatter(x=dv["timestamp"], y=dv["soil_moisture_pct"], name="Soil saturation", line=dict(color="#38BDF8")), secondary_y=True)
    fig_ari.update_layout(title="Antecedent rainfall vs soil saturation", height=300, margin=dict(l=10, r=10, t=45, b=10), legend=dict(orientation="h", y=-0.2), **DARK)
    c2.plotly_chart(fig_ari, **WIDE, key="ari")

    a1, a2 = st.columns(2)
    corr = dv[cols + ["threat_score"]].corr()
    fig_c = px.imshow(corr, text_auto=".2f", color_continuous_scale="RdBu_r", zmin=-1, zmax=1, title="Sensor correlation")
    fig_c.update_layout(height=380, margin=dict(l=10, r=10, t=50, b=10), **DARK)
    a1.plotly_chart(fig_c, **WIDE, key="corr")
    fig_s = px.scatter(dv.assign(size=dv["rainfall_pct"] + 2), x="soil_moisture_pct", y="tilt_angle_deg", color="threat_score",
                       size="size", color_continuous_scale="Turbo", title="Saturation vs tilt (size = rainfall)")
    fig_s.update_layout(height=380, margin=dict(l=10, r=10, t=50, b=10), **DARK)
    a2.plotly_chart(fig_s, **WIDE, key="scatter")
    st.plotly_chart(threat_timeline(dv), **WIDE, key="timeline")


def tab_forecast(S):
    st.subheader(f"Probabilistic forecast - next {FC_H} min")
    st.caption("Linear-trend Monte-Carlo (slope uncertainty + residual noise, 400 paths). Approximate and best for short horizons.")
    m = st.columns(4)
    pf = lambda p: "n/a" if p is None else f"{p * 100:.0f} %"
    m[0].metric("P(tilt ≥ critical)", pf(S.p_tilt))
    m[1].metric("P(saturation ≥ critical)", pf(S.p_moist))
    m[2].metric("P(FoS < 1)", pf(S.p_fos))
    m[3].metric("Tilt trend", f"{S.mc_tilt['slope'] * 60:+.2f} °/h" if S.mc_tilt else "n/a")
    c1, c2 = st.columns(2)
    c1.plotly_chart(fan_chart(S.df, "tilt_angle_deg", S.mc_tilt, TH["tilt_crit"], "Tilt forecast (°)"), **WIDE, key="fan_tilt")
    c2.plotly_chart(fan_chart(S.df, "soil_moisture_pct", S.mc_moist, TH["moist_crit"], "Soil saturation forecast (%)"), **WIDE, key="fan_moist")
    if S.mc_tilt is None:
        st.info("Need at least 8 samples of history for a forecast.")


def tab_geotech(S):
    s = S.sensors
    st.subheader("Geotechnical lab")
    s1, s2, s3, s4 = st.columns(4)
    sim_m = s1.slider("Soil saturation (%)", 0, 100, int(s["soil_moisture_pct"]), key="sim_m")
    sim_r = s2.slider("Rainfall (%)", 0, 100, int(s["rainfall_pct"]), key="sim_r")
    sim_t = s3.slider("Tilt (°)", 0.0, 45.0, float(min(s["tilt_angle_deg"], 45.0)), 0.5, key="sim_t")
    sim_v = s4.slider("Tremor pulses", 0, 10, int(min(s["vibration_pulses"], 10)), key="sim_v")
    sc = float(threat_score(sim_m, sim_r, sim_t, sim_v))
    sf = float(factor_of_safety(sim_t, sim_m, SOIL))
    stier = classify(sc, sim_t, sim_m, sim_r, sf, TH)
    srel = reliability(sim_t, sim_m, SOIL, CV)
    r = st.columns(4)
    r[0].metric("Scenario tier", stier["name"])
    r[1].metric("Threat score", f"{sc:.1f}", f"{sc - S.score:+.1f} vs live", delta_color="inverse")
    r[2].metric("FoS", f"{sf:.2f}", f"{sf - S.fos:+.2f} vs live")
    r[3].metric("P(failure)", f"{srel['pf'] * 100:.1f} %", f"{(srel['pf'] - S.rel['pf']) * 100:+.1f} pts", delta_color="inverse")
    st.plotly_chart(fos_heatmap(SOIL, (s["soil_moisture_pct"], s["tilt_angle_deg"]), (sim_m, sim_t)), **WIDE, key="fos_map")

    a, b = st.columns(2)
    a.plotly_chart(tornado(s, SOIL), **WIDE, key="tornado")
    b.plotly_chart(rel_hist(S.rel), **WIDE, key="relhist")

    st.markdown("#### Storm projection")
    p1, p2, p3 = st.columns(3)
    hours = p1.slider("Horizon (h)", 1, 24, 12, key="p_h")
    rise = p2.slider("Saturation rise (%/h)", 0.0, 10.0, 2.0, 0.5, key="p_rise")
    creep = p3.slider("Tilt creep (°/h)", 0.0, 2.0, 0.3, 0.05, key="p_creep")
    hh = np.arange(0, hours + 0.01, 0.25)
    pf_ = factor_of_safety(np.clip(s["tilt_angle_deg"] + creep * hh, 0, 45), np.clip(s["soil_moisture_pct"] + rise * hh, 0, 100), SOIL)
    fig_p = go.Figure(go.Scatter(x=hh, y=pf_, line=dict(color="#38BDF8", width=3), name="Projected FoS"))
    fig_p.add_hline(y=1.0, line_dash="dash", line_color="#EF4444", annotation_text="Failure (FoS = 1)")
    fig_p.update_layout(xaxis_title="Hours from now", yaxis_title="FoS", height=300, margin=dict(l=10, r=10, t=20, b=10), **DARK)
    st.plotly_chart(fig_p, **WIDE, key="proj")
    below = np.where(pf_ < 1.0)[0]
    if len(below):
        st.error(f"Projected slope failure in about {hh[below[0]]:.2f} h under this scenario.")
    else:
        st.success(f"Slope stays stable over {hours} h (minimum FoS {pf_.min():.2f}).")


def tab_replay(S):
    df = S.df
    n = len(df)
    st.subheader("Incident replay")
    ago = st.slider("Samples ago (0 = now)", 0, n - 1, 0, key="replay_ago")
    i = n - 1 - ago
    row = df.iloc[i]
    rs = {c: float(row[c]) for c in SENSOR_COLS}
    tname = row["risk_tier"]
    st.markdown(f"**{row['timestamp'].strftime('%Y-%m-%d %H:%M:%S')} UTC** &nbsp; "
                f"<span class='chip' style='background:{TIERS[tname]['color']};color:#0B1120'>{tname}</span>", unsafe_allow_html=True)
    m = st.columns(5)
    m[0].metric("Threat", f"{row['threat_score']:.1f}")
    m[1].metric("FoS", f"{row['fos']:.2f}")
    m[2].metric("Saturation", f"{rs['soil_moisture_pct']:.1f}%")
    m[3].metric("Rainfall", f"{rs['rainfall_pct']:.1f}%")
    m[4].metric("Tilt", f"{rs['tilt_angle_deg']:.1f}°")
    c1, c2 = st.columns([1, 1])
    c1.plotly_chart(make_3d(rs, "Elevation", TIERS[tname]["color"]), **WIDE, key="replay3d")
    c2.plotly_chart(trend_figure(df.tail(120), TH, 0, False, None, marker_ts=row["timestamp"]), **WIDE, key="replay_trend")


def tab_health(S):
    h = S.health
    k = st.columns(4)
    k[0].metric("Health score", f"{h['score']:.0f}/100", h["label"], delta_color="off")
    k[1].metric("Last record", "n/a" if h["age_s"] is None else f"{h['age_s']:.0f} s ago")
    k[2].metric("Sampling interval", "n/a" if not h["interval_s"] else f"{h['interval_s']:.0f} s")
    k[3].metric("Uplink", str(S.network.get("active_interface", "-")), f"RSSI {S.network.get('rssi', '-')} dBm", delta_color="off")
    icons = {"ok": "🟢", "warn": "🟡", "bad": "🔴"}
    show_df(pd.DataFrame([{"": icons[c["status"]], "Check": c["check"], "Detail": c["detail"]} for c in h["checks"]]))
    if S.extras:
        st.caption("Power telemetry: " + ", ".join(f"{k} = {v:g}" for k, v in S.extras.items()))



def architecture_figure():
    """Plotly architecture diagram; no additional visualization dependency needed."""
    fig = go.Figure()
    nodes = [
        {"x0": 0.15, "x1": 2.15, "y0": 2.55, "y1": 3.90, "title": "Sensor layer",
         "body": "Soil moisture • FC-37<br>DHT22 • SW-420 • MPU6050", "color": "#38BDF8"},
        {"x0": 2.55, "x1": 4.55, "y0": 2.55, "y1": 3.90, "title": "ESP32 edge",
         "body": "Read sensors • validate<br>Create telemetry JSON", "color": "#818CF8"},
        {"x0": 4.95, "x1": 6.95, "y0": 2.55, "y1": 3.90, "title": "Network uplink",
         "body": "Wi-Fi primary<br>SIM800L GPRS fallback", "color": "#34D399"},
        {"x0": 7.35, "x1": 9.35, "y0": 2.55, "y1": 3.90, "title": "Firebase RTDB",
         "body": "/nodes/{id}/live<br>/nodes/{id}/history", "color": "#FBBF24"},
        {"x0": 9.75, "x1": 11.85, "y0": 2.55, "y1": 3.90, "title": "Streamlit Hub",
         "body": "Risk score • FoS<br>Forecasts • health • exports", "color": "#F472B6"},
        {"x0": 2.65, "x1": 4.45, "y0": 0.45, "y1": 1.50, "title": "LittleFS offline queue",
         "body": "Store locally if both links fail", "color": "#94A3B8"},
    ]
    connectors = [
        (2.15, 3.22, 2.55, 3.22), (4.55, 3.22, 4.95, 3.22),
        (6.95, 3.22, 7.35, 3.22), (9.35, 3.22, 9.75, 3.22),
        (3.15, 2.55, 3.15, 1.50), (3.95, 1.50, 3.95, 2.55),
    ]
    for x0, y0, x1, y1 in connectors:
        fig.add_annotation(
            x=x1, y=y1, ax=x0, ay=y0, xref="x", yref="y", axref="x", ayref="y",
            text="", showarrow=True, arrowhead=2, arrowsize=1.1,
            arrowwidth=2, arrowcolor="#94A3B8",
        )
    for n in nodes:
        fig.add_shape(
            type="rect", x0=n["x0"], x1=n["x1"], y0=n["y0"], y1=n["y1"],
            line={"color": n["color"], "width": 2},
            fillcolor="rgba(15,23,42,0.96)", layer="below",
        )
        cx, cy = (n["x0"] + n["x1"]) / 2, (n["y0"] + n["y1"]) / 2
        fig.add_annotation(
            x=cx, y=cy + 0.20, text=f"<b>{n['title']}</b>", showarrow=False,
            font={"size": 12, "color": n["color"]}, xanchor="center", yanchor="middle",
        )
        fig.add_annotation(
            x=cx, y=cy - 0.19, text=n["body"], showarrow=False,
            font={"size": 10, "color": "#E2E8F0"}, xanchor="center", yanchor="middle",
        )
    fig.add_annotation(x=4.50, y=1.98, text="buffer / replay", showarrow=False,
                       font={"size": 9, "color": "#94A3B8"}, xanchor="left")
    fig.update_layout(
        height=455, margin={"l": 5, "r": 5, "t": 8, "b": 8},
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font={"color": "#E2E8F0"}, showlegend=False,
    )
    fig.update_xaxes(visible=False, range=[0, 12.0], fixedrange=True)
    fig.update_yaxes(visible=False, range=[0.1, 4.25], fixedrange=True)
    return fig


def tab_architecture(S):
    st.subheader("GeoAlert system architecture")
    st.caption("End-to-end sensor data path for the selected node.")
    st.plotly_chart(architecture_figure(), **WIDE, key="architecture_diagram")
    m = st.columns(4)
    m[0].metric("Telemetry", "LIVE" if S.is_live else "OFFLINE / DEMO")
    m[1].metric("Node ID", NODE_ID)
    m[2].metric("Active uplink", str(S.network.get("active_interface", "Unknown")))
    m[3].metric("History source", "Firebase" if S.hist_is_real else "Synthetic fallback")

    st.markdown("#### Data flow")
    st.markdown(
        "1. **Sensors → ESP32:** soil moisture, rain index, temperature, humidity, vibration and IMU readings are sampled and packaged as JSON.\\n"
        "2. **ESP32 → network:** Wi-Fi is attempted first; SIM800L GPRS is the fallback. If both links fail, telemetry is buffered in LittleFS and replayed after reconnection.\\n"
        "3. **Firebase Realtime Database:** live readings are written to /nodes/<node_id>/live; historical samples are written to /nodes/<node_id>/history.\\n"
        "4. **Streamlit Command Hub:** reads Firebase data and visualizes sensor status, hazard score, infinite-slope Factor of Safety, history, forecasts, anomalies and reports."
    )
    st.info(
        "This version has no cloud SMS worker and no physical siren/relay. "
        "Direct SIM800L SMS is an optional firmware-side feature and is disabled by default; "
        "webhook notifications can be configured separately."
    )


def tab_response(S):
    st.subheader("Alerts & event log")
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("#### Current assessment")
        st.metric("Hazard tier", S.tier["name"])
        st.metric("Threat score", f"{S.score:.1f}/100")
        st.metric("Factor of Safety", f"{S.fos:.2f}")
    with col2:
        st.markdown("#### Notification configuration")
        if NOTIFY_ON and WEBHOOK:
            st.success("Webhook escalation notifications are enabled.")
        else:
            st.caption("Webhook escalation notifications are off. Configure them from the sidebar if needed.")
        st.info(
            "This dashboard records and displays events; it does not dispatch SMS. "
            "Optional direct SIM800L SMS is configured in the firmware and is disabled by default. "
            "No physical siren or relay is installed."
        )

    st.markdown("#### Session alert log")
    log = st.session_state.alert_log
    unacked = sum(1 for a in log if not a.get("ack"))
    c1, c2 = st.columns([1, 3])
    c1.metric("Unacknowledged", unacked)
    if c2.button("✅ Acknowledge all", key="btn_ack"):
        for a in log:
            a["ack"] = True
        st.rerun(scope="fragment") if NEW_API else None
    if log:
        show_df(pd.DataFrame(log))
    else:
        st.caption("No events yet.")



def tab_export(S):
    st.subheader("Reports & data export")
    st.write("Build once, then download individual formats or the full ZIP bundle. Files persist while auto-refresh continues.")
    st.markdown(f"<div class='brief'>{S.narrative}</div>", unsafe_allow_html=True)
    if st.button("⚙️ Build export package", type="primary", key="btn_build"):
        with st.spinner("Generating Excel, PDF, HTML, GeoJSON, JSON, CSV and ZIP..."):
            st.session_state.exports = build_package(S)
    ex = st.session_state.exports
    if not ex:
        return
    f, stamp = ex["files"], ex["stamp"]
    st.success(f"Package built {ex['built_utc']} UTC")
    base = f"GeoAlert_{NODE_ID}_{stamp}"
    d = st.columns(4)
    d[0].download_button("📦 Full bundle (.zip)", f["zip"], f"{base}.zip", "application/zip", type="primary", **WIDE)
    d[1].download_button("📊 Excel workbook", f["xlsx"], f"{base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", **WIDE)
    d[2].download_button("🌐 Interactive HTML", f["html"], f"{base}.html", "text/html", **WIDE)
    if f["pdf"]:
        d[3].download_button("📕 PDF report", f["pdf"], f"{base}.pdf", "application/pdf", **WIDE)
    else:
        d[3].caption("PDF needs `pip install reportlab matplotlib`")
    e = st.columns(4)
    e[0].download_button("📄 CSV", f["csv"], f"{base}.csv", "text/csv", **WIDE)
    e[1].download_button("🧾 JSON snapshot", f["json"], f"{base}.json", "application/json", **WIDE)
    e[2].download_button("🗺️ GeoJSON (GIS)", f["geojson"], f"{base}.geojson", "application/geo+json", **WIDE)


# ==============================================================================
# 10. MAIN RENDER
# ==============================================================================
def render():
    S = build_context()
    log_and_notify(S)
    c = S.tier["color"]
    pulse = "pulse" if S.tier["level"] >= 2 else ""
    raw_chip = f"<span class='chip'>raw signal: {S.raw_tier['name']}</span>" if S.pending else ""
    fleet_chip = f"<span class='chip'>{len(S.fleet)} node(s)</span>"
    st.markdown(
        f"""
    <div class="status-card" style="background:linear-gradient(90deg,#0F172A 0%,#1E293B 100%);border-left:8px solid {c};box-shadow:0 0 24px {c}44;">
      <div style="display:flex;justify-content:space-between;align-items:center;flex-wrap:wrap;gap:12px;">
        <div>
          <h1 style="margin:0;font-size:1.9rem;color:#fff;font-weight:800;">GEOALERT DISASTER INTELLIGENCE</h1>
          <p style="margin:6px 0 0 0;color:#94A3B8;font-size:.92rem;">
            NODE <span style="color:#38BDF8;font-weight:600;">{NODE_ID}</span> &nbsp;|&nbsp;
            UPLINK <span style="color:#4ADE80;font-weight:600;">{S.network.get('active_interface', 'ONLINE')}</span> (RSSI {S.network.get('rssi', '-')} dBm) &nbsp;|&nbsp;
            <span style="color:{'#4ADE80' if S.is_live else '#FACC15'};">{'LIVE STREAMING' if S.is_live else 'OFFLINE / DEMO'}</span>
          </p>
          <p style="margin:8px 0 0 0;">{raw_chip}{fleet_chip}<span class="chip">health {S.health['label']} {S.health['score']:.0f}</span>
             <span class="chip">history: {'device log' if S.hist_is_real else 'synthetic'}</span></p>
        </div>
        <div style="text-align:right;">
          <div style="font-size:.8rem;color:#94A3B8;text-transform:uppercase;">Hazard risk level</div>
          <div class="{pulse}" style="font-size:2.1rem;font-weight:900;color:{c};">{S.tier['name']}</div>
        </div>
      </div>
    </div>""",
        unsafe_allow_html=True,
    )
    if st.session_state.get("auth_error"):
        st.warning(f"Firebase authentication failed: {st.session_state['auth_error']}")
    if S.tier["level"] >= 2:
        acts = "\n".join(f"- {a}" for a in PROTOCOL[S.tier["name"]])
        st.error(f"⚠️ **{S.tier['name']}** - threat **{S.score:.1f}/100**, FoS **{S.fos:.2f}**, probability of failure **{S.rel['pf'] * 100:.1f}%**. "
                 f"Tilt hits critical in **{fmt_eta(S.fc_tilt)}**, saturation in **{fmt_eta(S.fc_moist)}**.\n\n{acts}")

    s, df = S.sensors, S.df

    def dl(col, unit=""):
        return f"{s[col] - df[col].iloc[-11]:+.1f}{unit}" if len(df) > 10 else None

    m = st.columns(7)
    m[0].metric("Soil Saturation", f"{s['soil_moisture_pct']:.1f}%", dl("soil_moisture_pct", "%"), delta_color="inverse")
    m[1].metric("Precipitation", f"{s['rainfall_pct']:.1f}%", dl("rainfall_pct", "%"), delta_color="inverse")
    m[2].metric("Slope Tilt", f"{s['tilt_angle_deg']:.1f}°", dl("tilt_angle_deg", "°"), delta_color="inverse")
    m[3].metric("Seismic Tremors", f"{s['vibration_pulses']} /s")
    m[4].metric("Temperature", f"{s['temperature_c']:.1f} °C")
    m[5].metric("Air Humidity", f"{s['humidity_pct']:.0f}%")
    m[6].metric("Pitch / Roll", f"{s['pitch_deg']:.1f}° / {s['roll_deg']:.1f}°")

    tabs = st.tabs(["🎯 Situation", "🗺️ Fleet", "🌐 3D Twin", "🏗️ Architecture", "📈 Analytics", "🔮 Forecast",
                    "🧪 Geotech Lab", "⏪ Replay", "🩺 Node Health", "🚨 Response", "📥 Reports & Export"])
    for tab, fn in zip(tabs, (tab_situation, tab_fleet, tab_twin, tab_architecture, tab_analytics, tab_forecast,
                              tab_geotech, tab_replay, tab_health, tab_response, tab_export)):
        with tab:
            fn(S)


if hasattr(st, "fragment"):
    st.fragment(run_every=refresh_s if auto_refresh else None)(render)()
else:  # Streamlit < 1.37 fallback
    render()
    if auto_refresh:
        time.sleep(refresh_s)
        st.rerun()